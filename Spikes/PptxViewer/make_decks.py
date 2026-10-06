"""Test decks for the PowerPoint viewer spike (ENH-12). Needs python-pptx and Pillow.

    python3 make_decks.py <out-dir>

Writes basics.pptx (title, bullets, theme fonts, notes), data.pptx (a table, a bar chart, a pie
chart) and shapes.pptx (images, a group inside a group, connectors, autoshapes, rotated text).
"""
import sys
from pathlib import Path

from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE, XL_LEGEND_POSITION
from pptx.enum.shapes import MSO_CONNECTOR, MSO_SHAPE
from pptx.util import Emu, Inches, Pt

out = Path(sys.argv[1] if len(sys.argv) > 1 else ".")
out.mkdir(parents=True, exist_ok=True)


def wide():
    p = Presentation()
    p.slide_width, p.slide_height = Inches(13.333), Inches(7.5)
    return p


def basics():
    p = wide()
    s = p.slides.add_slide(p.slide_layouts[0])
    s.shapes.title.text = "Quarterly review"
    s.placeholders[1].text = "Theme fonts, placeholders and notes"
    s.notes_slide.notes_text_frame.text = "Speaker notes for the title slide."

    s = p.slides.add_slide(p.slide_layouts[1])
    s.shapes.title.text = "What changed"
    tf = s.placeholders[1].text_frame
    tf.text = "Revenue up 12% on the quarter"
    for text, level in [("New customers: 41", 1), ("Churn down to 2.1%", 1), ("Hiring paused", 0), ("Two roles still open", 1)]:
        para = tf.add_paragraph()
        para.text, para.level = text, level
    s.notes_slide.notes_text_frame.text = "Mention the churn number first."

    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Mixed runs in one box"
    box = s.shapes.add_textbox(Inches(1), Inches(2), Inches(11), Inches(3))
    para = box.text_frame.paragraphs[0]
    for text, bold, italic, color, size in [("Bold ", True, False, None, 28), ("italic ", False, True, None, 28),
                                            ("red ", False, False, RGBColor(0xC0, 0x20, 0x20), 28), ("small", False, False, None, 14)]:
        r = para.add_run()
        r.text = text
        r.font.bold, r.font.italic, r.font.size = bold, italic, Pt(size)
        if color: r.font.color.rgb = color
    box.text_frame.add_paragraph().text = "A second paragraph in the theme's body font."
    p.save(out / "basics.pptx")


def data():
    p = wide()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "A table"
    rows, cols = 4, 3
    t = s.shapes.add_table(rows, cols, Inches(1), Inches(1.8), Inches(11), Inches(3)).table
    for c, h in enumerate(["Region", "Q2", "Q3"]): t.cell(0, c).text = h
    for r, row in enumerate([["North", "1.2m", "1.4m"], ["South", "0.9m", "1.1m"], ["West", "2.0m", "1.8m"]], start=1):
        for c, v in enumerate(row): t.cell(r, c).text = v
    t.cell(3, 0).merge(t.cell(3, 0))

    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "A bar chart"
    cd = CategoryChartData()
    cd.categories = ["North", "South", "West"]
    cd.add_series("Q2", (1.2, 0.9, 2.0))
    cd.add_series("Q3", (1.4, 1.1, 1.8))
    ch = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(1), Inches(1.6), Inches(11), Inches(5.4), cd).chart
    ch.has_legend, ch.legend.position, ch.legend.include_in_layout = True, XL_LEGEND_POSITION.BOTTOM, False

    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "A pie chart"
    cd = CategoryChartData()
    cd.categories = ["Product", "Services", "Support"]
    cd.add_series("Mix", (0.55, 0.3, 0.15))
    ch = s.shapes.add_chart(XL_CHART_TYPE.PIE, Inches(3.5), Inches(1.6), Inches(6), Inches(5.4), cd).chart
    ch.has_legend = True
    p.save(out / "data.pptx")


def shapes():
    img = out / "_photo.png"
    im = Image.new("RGB", (640, 400), (40, 90, 160))
    d = ImageDraw.Draw(im)
    for i in range(0, 640, 40): d.rectangle([i, 0, i + 20, 400], fill=(230, 160, 60))
    d.ellipse([220, 100, 420, 300], fill=(250, 250, 250))
    im.save(img)

    p = wide()
    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Images"
    s.shapes.add_picture(str(img), Inches(1), Inches(1.8), Inches(5))
    pic = s.shapes.add_picture(str(img), Inches(7), Inches(1.8), Inches(5))
    pic.rotation = 8

    s = p.slides.add_slide(p.slide_layouts[5])
    s.shapes.title.text = "Groups and connectors"
    outer = s.shapes.add_group_shape()
    a = outer.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1), Inches(2), Inches(3), Inches(1.2))
    a.text = "Plan"
    inner = outer.shapes.add_group_shape()
    b = inner.shapes.add_shape(MSO_SHAPE.OVAL, Inches(5.5), Inches(2), Inches(2.4), Inches(1.2))
    b.text = "Build"
    c = inner.shapes.add_shape(MSO_SHAPE.DIAMOND, Inches(9.5), Inches(1.8), Inches(2.4), Inches(1.6))
    c.text = "Ship?"
    for x1, x2 in [(4, 5.5), (7.9, 9.5)]:
        ln = s.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, Inches(x1), Inches(2.6), Inches(x2), Inches(2.6))
        ln.line.width = Pt(2)
    star = s.shapes.add_shape(MSO_SHAPE.STAR_5_POINT, Inches(1), Inches(4.5), Inches(2), Inches(2))
    star.fill.solid(); star.fill.fore_color.rgb = RGBColor(0xF0, 0xB0, 0x20)
    rot = s.shapes.add_textbox(Inches(5), Inches(5), Inches(5), Inches(1))
    rot.text_frame.text = "Rotated text box"
    rot.rotation = -10
    p.save(out / "shapes.pptx")
    img.unlink()


basics(); data(); shapes()
print("wrote", *(out / n for n in ["basics.pptx", "data.pptx", "shapes.pptx"]))
