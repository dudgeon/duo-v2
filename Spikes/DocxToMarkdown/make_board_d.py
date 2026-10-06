"""Board D's Report.docx (DL-123): inferred headings, typed bullets, two tracked changes, a comment and three
pictures, so the copy's notice shows its summary line. Synthetic. python3 make_board_d.py <out.docx>"""
import io, struct, sys, zlib
from docx import Document
from docx.oxml import parse_xml
from docx.oxml.ns import nsdecls
from docx.shared import Pt, Inches
def png(w,h,rgb):
    raw=b"".join(b"\x00"+bytes(rgb)*w for _ in range(h))
    c=lambda t,d: struct.pack(">I",len(d))+t+d+struct.pack(">I",zlib.crc32(t+d)&0xFFFFFFFF)
    return b"\x89PNG\r\n\x1a\n"+c(b"IHDR",struct.pack(">IIBBBBB",w,h,8,2,0,0,0))+c(b"IDAT",zlib.compress(raw))+c(b"IEND",b"")
d=Document(); d.styles["Normal"].font.size=Pt(11)
def big(t,size):
    r=d.add_paragraph().add_run(t); r.bold=True; r.font.size=Pt(size)
big("Quarterly Garden Report",24)
d.add_paragraph("This report covers the imaginary community garden on Example Street.")
big("Summary",16)
p=d.add_paragraph("Yields were ")
p._p.append(parse_xml(f'<w:ins {nsdecls("w")} w:id="1" w:author="Reviewer A" w:date="2026-01-01T00:00:00Z"><w:r><w:t xml:space="preserve">well </w:t></w:r></w:ins>'))
p.add_run("up this season, and the tomatoes did best.")
p._p.append(parse_xml(f'<w:del {nsdecls("w")} w:id="2" w:author="Reviewer B" w:date="2026-01-02T00:00:00Z"><w:r><w:delText xml:space="preserve"> Again.</w:delText></w:r></w:del>'))
big("What we planted",13)
for t in ["• Tomatoes","• Beans"]: d.add_paragraph(t)
r=d.add_paragraph().add_run("We will spend fifty units on seeds.")
d.add_comment(r,text="Is fifty enough?",author="Reviewer A",initials="RA")
big("Harvest",13)
for rgb in [(60,140,80),(200,80,60),(80,100,200)]:
    d.add_picture(io.BytesIO(png(60,30,rgb)),width=Inches(1.5))
d.save(sys.argv[1])
