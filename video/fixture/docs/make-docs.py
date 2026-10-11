#!/usr/bin/env python3
"""Builds docs/exec-review-oct14.pptx and docs/interview-synthesis.docx from content.json (sample data only).
Run:  uv run --with python-pptx --with python-docx video/fixture/docs/make-docs.py"""
import json, pathlib
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.chart.data import CategoryChartData
from pptx.enum.chart import XL_CHART_TYPE, XL_LABEL_POSITION
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE
from docx import Document
from docx.shared import Pt as DPt, RGBColor as DRGB, Inches as DIn
from docx.enum.text import WD_ALIGN_PARAGRAPH

here = pathlib.Path(__file__).parent
c = json.load(open(here / "content.json"))
INK, MUTED, ACC, BG, LINE = RGBColor(0x1A,0x1D,0x21), RGBColor(0x6B,0x72,0x80), RGBColor(0x1F,0x4D,0x3A), RGBColor(0xF7,0xF6,0xF3), RGBColor(0xE3,0xE5,0xE8)
FONT = "Helvetica Neue"

def text(slide, x, y, w, h, s, size, color=INK, bold=False, align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf = tb.text_frame; tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]; p.alignment = align; r = p.add_run(); r.text = s
    r.font.size, r.font.bold, r.font.name = Pt(size), bold, FONT; r.font.color.rgb = color
    return tb

def base(prs, title, n):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    s.background.fill.solid(); s.background.fill.fore_color.rgb = BG
    bar = s.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(0.22), prs.slide_height); bar.fill.solid(); bar.fill.fore_color.rgb = ACC; bar.line.fill.background()
    text(s, 0.8, 0.55, 11.7, 1.2, title, 34, INK, True, anchor=MSO_ANCHOR.MIDDLE)
    text(s, 0.8, 7.0, 8, 0.3, "Northwind Supply · checkout-redesign · exec review", 11, MUTED)
    text(s, 11.9, 7.0, 0.8, 0.3, str(n), 11, MUTED, align=PP_ALIGN.RIGHT)
    return s

prs = Presentation(); prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
for n, sl in enumerate(c["deck"]["slides"], 1):
    if sl["kind"] == "title":
        s = prs.slides.add_slide(prs.slide_layouts[6]); s.background.fill.solid(); s.background.fill.fore_color.rgb = ACC
        text(s, 0.9, 2.5, 11.5, 1.6, sl["title"], 50, RGBColor(255,255,255), True, anchor=MSO_ANCHOR.BOTTOM)
        text(s, 0.9, 4.3, 11.5, 0.6, sl["subtitle"], 22, RGBColor(0xC9,0xDD,0xD3))
        continue
    s = base(prs, sl["title"], n)
    if sl["kind"] == "chart":
        ch = sl["chart"]; cd = CategoryChartData(); cd.categories = [b[0] for b in ch["bars"]]; cd.add_series(ch["label"], [b[1] for b in ch["bars"]])
        gf = s.shapes.add_chart(XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.8), Inches(1.9), Inches(8.2), Inches(4.8), cd); chart = gf.chart
        chart.has_legend = False; chart.has_title = True; chart.chart_title.text_frame.text = ch["label"]
        chart.chart_title.text_frame.paragraphs[0].runs[0].font.size = Pt(14); chart.chart_title.text_frame.paragraphs[0].runs[0].font.color.rgb = MUTED
        pl = chart.plots[0]; pl.gap_width = 60; pl.has_data_labels = True
        pl.data_labels.font.size = Pt(16); pl.data_labels.font.bold = True; pl.data_labels.position = XL_LABEL_POSITION.OUTSIDE_END
        ser = pl.series[0]; peak = max(b[1] for b in ch["bars"])
        for i, b in enumerate(ch["bars"]):
            pt = ser.points[i]; pt.format.fill.solid(); pt.format.fill.fore_color.rgb = ACC if b[1] == peak else RGBColor(0xB9,0xC4,0xBE)
        chart.value_axis.visible = False; chart.value_axis.has_major_gridlines = False
        chart.category_axis.tick_labels.font.size = Pt(15); chart.category_axis.format.line.color.rgb = LINE
        card = s.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(9.4), Inches(2.4), Inches(3.3), Inches(2.6)); card.fill.solid(); card.fill.fore_color.rgb = RGBColor(255,255,255); card.line.color.rgb = LINE
        text(s, 9.6, 2.55, 2.9, 1.0, "38%", 54, ACC, True)
        text(s, 9.6, 3.7, 2.9, 1.2, sl["note"][4:] if sl["note"].startswith("38% ") else sl["note"], 17, INK)
    else:
        tb = s.shapes.add_textbox(Inches(0.9), Inches(2.1), Inches(11.4), Inches(4.4)); tf = tb.text_frame; tf.word_wrap = True
        for i, b in enumerate(sl["bullets"]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph(); p.space_after = Pt(22)
            r = p.add_run(); r.text = "•  " + b; r.font.size, r.font.name = Pt(30), FONT; r.font.color.rgb = INK
prs.save(here / c["deck"]["file"])

d = c["doc"]; doc = Document()
sec = doc.sections[0]; sec.left_margin = sec.right_margin = DIn(1.1); sec.top_margin = sec.bottom_margin = DIn(1)
st = doc.styles
from docx.oxml.ns import qn
def style(name, size, color=None, bold=None, italic=None):
    f = st[name].font; f.name = "Arial"; f.size = DPt(size)
    rpr = st[name].element.get_or_add_rPr(); rf = rpr.find(qn("w:rFonts"))
    for a in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
        if rf is not None and rf.get(qn(a)) is not None: del rf.attrib[qn(a)]
    if bold is not None: f.bold = bold
    if italic is not None: f.italic = italic
    if color:
        f.color.rgb = DRGB(*color)
        c = rpr.find(qn("w:color"))
        for a in ("w:themeColor", "w:themeShade", "w:themeTint"):
            if c.get(qn(a)) is not None: del c.attrib[qn(a)]
style("Title Char", 24, (0x1F,0x4D,0x3A), True)   # the renderer reads the linked character styles too
style("Heading 1 Char", 16, (0x1F,0x4D,0x3A), True)
style("Normal", 11)
style("Title", 24, (0x1F,0x4D,0x3A), True)
style("Heading 1", 16, (0x1F,0x4D,0x3A), True)
style("List Bullet", 11)
style("Quote", 11.5, (0x37,0x41,0x51), None, True)
st["Heading 1"].paragraph_format.space_before = DPt(18)
doc.add_paragraph(d["title"], style="Title")
p = doc.add_paragraph(d["byline"]); p.runs[0].font.color.rgb = DRGB(0x6B,0x72,0x80)
for s_ in d["sections"]:
    doc.add_heading(s_["heading"], level=1)
    if "body" in s_: doc.add_paragraph(s_["body"])
    for q in s_.get("quotes", []):
        qp = doc.add_paragraph(q, style="Quote")
    for b in s_.get("bullets", []): doc.add_paragraph(b, style="List Bullet")
doc.save(here / d["file"])
print("ok")
